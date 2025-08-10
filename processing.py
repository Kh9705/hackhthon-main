# processing.py
import httpx, fitz, numpy as np, faiss, time, io, re, zipfile, magic, docx, extract_msg, pandas as pd
from PIL import Image
import pytesseract
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from fastapi import HTTPException
from langchain.text_splitter import RecursiveCharacterTextSplitter

from config import ENCODER, CHUNK_CONFIG, SMALL_DOC_THRESHOLD, LARGE_DOC_THRESHOLD

magic_obj = magic.Magic(mime=True)

def _extract_text_from_file(file_extension: str, doc_bytes: bytes) -> str:
    # This function would contain all the if/elif logic for PDF, DOCX, PPTX, etc.
    # To save space, the full implementation is omitted here, but you would paste
    # your original text extraction logic into this helper function.
    full_text = ""
    if file_extension == ".pdf":
        with fitz.open("pdf", doc_bytes) as doc:
            full_text = "".join(page.get_text() for page in doc)
    elif file_extension == ".docx":
        doc_stream = io.BytesIO(doc_bytes)
        document = docx.Document(doc_stream)
        full_text = "\n".join([para.text for para in document.paragraphs])
    # ... and so on for all other file types.
    else:
        raise HTTPException(status_code=415, detail=f"Unsupported file type: '{file_extension}'.")
    return full_text


def process_and_index_document(doc_url: str, document_cache: dict):
    """Downloads, processes, and indexes a document, storing it in the provided cache."""
    print(f"[INFO] Processing document: {doc_url}")
    try:
        # 1. Download & Determine File Type
        response = httpx.get(doc_url, timeout=120.0)
        response.raise_for_status()
        doc_bytes = response.content
        
        match = re.search(r"\.(pdf|docx|eml|...|htm)(?=\?|$)", doc_url, re.IGNORECASE)
        file_extension = match.group(0).lower() if match else f".{magic_obj.from_buffer(doc_bytes).split('/')[-1]}"

        # 2. Extract Text
        full_text = _extract_text_from_file(file_extension, doc_bytes)
        print(f"[INFO] Extracted {len(full_text)} characters.")

        # 3. Dynamic Chunking
        doc_len = len(full_text)
        size_category = "medium"
        if doc_len < SMALL_DOC_THRESHOLD: size_category = "small"
        elif doc_len > LARGE_DOC_THRESHOLD: size_category = "large"
        chunk_size, chunk_overlap, k_value = CHUNK_CONFIG[size_category]
        
        text_splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap)
        chunks = text_splitter.split_text(full_text)

        # 4. Embed and Index
        vectors = ENCODER.encode(chunks, show_progress_bar=True).astype("float32")
        index = faiss.IndexFlatL2(vectors.shape[1])
        index.add(vectors)

        # 5. Store in cache (passed as an argument)
        document_cache[doc_url] = {
            "index": index, "chunks": chunks, "timestamp": time.time(), "k_value": k_value
        }
        print(f"[SUCCESS] Indexed: {doc_url}")

    except Exception as e:
        print(f"[ERROR] Document processing failed: {e}")
        if isinstance(e, HTTPException): raise
        raise HTTPException(status_code=500, detail=f"Failed to process document: {e}")